from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "word_assets"
OUTPUT = ROOT / "初三升高一物理竞赛测试题_填空题_含答案.docx"

# Design system: compact_reference_guide with a named "A4 exam form" override.
# Page: A4, 1.55 cm side margins, 1.45 cm top/bottom.
# Body: Arial Unicode MS 10.5 pt, 1.18 line spacing, 2.5 pt after.
# Hierarchy: centered 21/13 pt black title stack; 11 pt bold question labels.
# Color: black/white/gray only, matching the supplied LaTeX test.


def set_run_font(run, name="Arial Unicode MS", size=10.5, bold=None, italic=None, color="000000"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def set_style_font(style, name, size, bold=False, color="000000"):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=0):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    total = sum(widths_dxa)

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideV", "insideH"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), "777777")

    old_grid = table._tbl.tblGrid
    for child in list(old_grid):
        old_grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        old_grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_bottom_border(paragraph, color="000000", size="6", space="2"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def shade_paragraph(paragraph, fill="F0F0F0"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    pbdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:space"), "4")
        edge.set(qn("w:color"), "777777")
        pbdr.append(edge)
    ppr.append(pbdr)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, "Arial Unicode MS", 9)


def add_text(doc, text, *, style=None, align=None, before=0, after=2.5, line=1.18,
             font="Arial Unicode MS", size=10.5, bold=False, keep_next=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep_next
    r = p.add_run(text)
    set_run_font(r, font, size, bold=bold)
    return p


def add_question_heading(doc, number, title):
    p = doc.add_paragraph(style="Question Heading")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{number}.  {title}")
    set_run_font(r, "Arial Unicode MS", 11, bold=True)
    return p


def add_question_text(doc, text):
    p = add_text(doc, text, after=2, line=1.18, keep_next=True)
    p.paragraph_format.first_line_indent = Cm(0.42)
    return p


def add_figure(doc, filename, width_cm, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    inline = run.add_picture(str(ASSETS / filename), width=Cm(width_cm))
    inline._inline.docPr.set("descr", alt)
    return p


def add_answer_heading(doc, number, answer):
    p = doc.add_paragraph(style="Answer Heading")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{number}.  {answer}")
    set_run_font(r, "Arial Unicode MS", 10.5, bold=True)
    return p


def add_answer_text(doc, text, *, indent=True, keep_next=False):
    p = add_text(doc, text, size=9.7, line=1.13, after=1.5, keep_next=keep_next)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.42)
    return p


def add_equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, "Cambria Math", 10.2)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.45)
section.bottom_margin = Cm(1.45)
section.left_margin = Cm(1.55)
section.right_margin = Cm(1.55)
section.header_distance = Cm(0.65)
section.footer_distance = Cm(0.65)

styles = doc.styles
normal = styles["Normal"]
set_style_font(normal, "Arial Unicode MS", 10.5)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(2.5)
normal.paragraph_format.line_spacing = 1.18

for name, font, size, bold in (
    ("Question Heading", "Arial Unicode MS", 11, True),
    ("Answer Heading", "Arial Unicode MS", 10.5, True),
):
    if name not in styles:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles[name]
    set_style_font(style, font, size, bold=bold)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(5)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.line_spacing = 1.1

# Running header/footer.
header = section.header
hp = header.paragraphs[0]
hp.text = ""
hp.paragraph_format.space_after = Pt(2)
hp.paragraph_format.tab_stops.add_tab_stop(Cm(17.8), WD_TAB_ALIGNMENT.RIGHT)
hr1 = hp.add_run("初三升高一物理竞赛测试")
set_run_font(hr1, "Arial Unicode MS", 8.8)
hr2 = hp.add_run("\t填空题专题")
set_run_font(hr2, "Arial Unicode MS", 8.8)
add_bottom_border(hp, size="4", space="1")

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_before = Pt(2)
add_page_field(fp)

# Title block.
title = add_text(doc, "初三升高一物理竞赛测试题", align=WD_ALIGN_PARAGRAPH.CENTER,
                 after=1, line=1.0, font="Arial Unicode MS", size=21, bold=True)
subtitle = add_text(doc, "——填空题（共 5 题）", align=WD_ALIGN_PARAGRAPH.CENTER,
                    after=6, line=1.0, font="Arial Unicode MS", size=13, bold=True)

info = doc.add_table(rows=1, cols=4)
set_table_geometry(info, [1500, 2600, 1500, 2600], indent_dxa=950)
for idx, text_value in enumerate(("姓名", "", "得分", "")):
    cell = info.cell(0, idx)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text_value)
    set_run_font(r, "Arial Unicode MS", 10.5)

spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(1)

note = doc.add_paragraph()
note.paragraph_format.left_indent = Cm(0.2)
note.paragraph_format.right_indent = Cm(0.2)
note.paragraph_format.space_before = Pt(2)
note.paragraph_format.space_after = Pt(4)
note.paragraph_format.line_spacing = 1.1
shade_paragraph(note, "F1F1F1")
r = note.add_run("说明：")
set_run_font(r, "Arial Unicode MS", 9.5, bold=True)
r = note.add_run("建议用时 45 分钟，满分 50 分。每题 10 分，每个空 5 分。除特别说明外，忽略空气阻力、导线电阻及电源内阻，结果可保留分数或根式。")
set_run_font(r, "Arial Unicode MS", 9.5)

# Questions, page 1.
add_question_heading(doc, 1, "运动图像")
add_question_text(doc, "甲、乙两物体在同一直线上从同一位置同时出发，运动方向始终相同。甲的速度随时间变化关系如图中实线所示；乙始终以 5 m/s 做匀速直线运动（图中虚线）。甲在前 12 s 内通过的路程为 __________ m；除出发时刻外，甲第一次追上乙的时刻为 t = __________ s。")
add_figure(doc, "fig1_motion.png", 12.2, "甲、乙两物体的速度—时间图像")

add_question_heading(doc, 2, "杠杆平衡")
add_question_text(doc, "如图，一根长 1.2 m、重 12 N 的均匀直杆 OB 可绕端点 O 自由转动。杆保持水平，在距 O 点 0.20 m 处悬挂一个重 30 N 的物体。若在 B 点沿竖直方向向上拉杆，平衡时拉力大小为 __________ N；若改为在 B 点沿与杆成 30° 角的方向斜向上拉杆，平衡时拉力大小为 __________ N。")
add_figure(doc, "fig2_lever.png", 12.3, "均匀杠杆受力示意图")

doc.add_page_break()

# Questions, page 2.
add_question_heading(doc, 3, "浮力与密度")
add_question_text(doc, "底面积为 50 cm²、高为 20 cm 的均匀圆柱体，密度为 0.80 g/cm³。它竖直静止在互不相溶的油和水中，初始时圆柱顶面恰与油面相平，如图所示。已知油的密度为 0.60 g/cm³，水的密度为 1.00 g/cm³。圆柱浸入水中的高度为 __________ cm；若在圆柱顶面中央轻放一个体积可忽略的小物块，重新平衡后圆柱下沉了 2.0 cm，则小物块的质量为 __________ g。")
add_figure(doc, "fig3_buoyancy.png", 6.2, "圆柱体竖直漂浮在油水分层中的示意图")

add_question_heading(doc, 4, "对称电路")
add_question_text(doc, "如图，菱形网络的五个电阻均为 6.0 Ω，理想电源电压为 12 V。接通电路后，流过电源的总电流为 __________ A；流过中间电阻 R₅ 的电流为 __________ A。")
add_figure(doc, "fig4_circuit.png", 6.2, "由五个相同电阻构成的对称桥式电路")

add_question_heading(doc, 5, "凸透镜成像")
add_question_text(doc, "物体和光屏固定在光具座上，两者间距为 100 cm，物体高度为 2.0 cm。在它们之间移动同一凸透镜，先后在相距 60 cm 的两个位置上，都能在光屏上得到清晰的像。该凸透镜的焦距为 __________ cm；两次成像中，较大像的高度为 __________ cm。")
add_figure(doc, "fig5_lens.png", 11.5, "物体、两个凸透镜位置与光屏的相对位置")

# Teacher answer key.
answer_title = add_text(doc, "参考答案与简析（教师用）", align=WD_ALIGN_PARAGRAPH.CENTER,
                        after=7, line=1.0, font="Arial Unicode MS", size=16, bold=True)
answer_title.paragraph_format.page_break_before = True

add_answer_heading(doc, 1, "72 m；16/3 s。")
add_answer_text(doc, "甲的路程等于 v–t 图像下的面积：", keep_next=True)
add_equation(doc, "s甲 = 1/2 × 4 × 8 + 6 × 8 + 1/2 × 2 × 8 = 72 m。")
add_answer_text(doc, "前 4 s 内甲通过 16 m，乙通过 20 m，甲尚未追上乙。随后甲做匀速运动，令 16 + 8(t − 4) = 5t，得 t = 16/3 s。")

add_answer_heading(doc, 2, "11 N；22 N。")
add_answer_text(doc, "对 O 点取力矩。竖直拉时：", keep_next=True)
add_equation(doc, "F₁ × 1.2 = 12 × 0.6 + 30 × 0.2，")
add_answer_text(doc, "得 F₁ = 11 N。斜拉时只有垂直于杆的分力产生力矩，F₂ sin 30° = F₁，故 F₂ = 22 N。")

add_answer_heading(doc, 3, "10 cm；40 g。")
add_answer_text(doc, "设圆柱浸入水中的高度为 h。由“排开液体的质量等于物体质量”：", keep_next=True)
add_equation(doc, "1.00 × 50h + 0.60 × 50(20 − h) = 0.80 × 50 × 20，")
add_answer_text(doc, "解得 h = 10 cm。下沉 2.0 cm 后，有 2.0 cm 高的一段由排开油变为排开水，增加的浮力等效质量为 (1.00 − 0.60) × 50 × 2.0 = 40 g。")

add_answer_heading(doc, 4, "2.0 A；0 A。")
add_answer_text(doc, "上、下两支路完全对称，C、D 两点等电势，所以 R₅ 中无电流。每条支路电阻均为 12 Ω，两支路并联后的等效电阻为 6 Ω，故电源电流 I = 12/6 = 2.0 A。")

add_answer_heading(doc, 5, "16 cm；8.0 cm。")
add_answer_text(doc, "设物体在 x = 0，光屏在 x = 100 cm。透镜位于第一个清晰位置 L₁ 时，物距 u = x₁，像距 v = 100 − x₁，因此 u + v = 100 cm。", keep_next=True)
add_answer_text(doc, "透镜公式 1/f = 1/u + 1/v 对 u、v 完全对称，所以同一组物、屏位置还有一个清晰位置，其物距和像距恰好互换：u′ = v，v′ = u。于是第二个透镜位置的坐标 x₂ = u′ = v。", keep_next=True)
add_equation(doc, "L₁L₂ = |x₂ − x₁| = |v − u| = 60 cm。")
add_answer_text(doc, "所以这里不是先假定“物距和像距之差是 60 cm”，而是由两个透镜位置的间距等于 |v − u| 得到。联立 u + v = 100 cm、|v − u| = 60 cm，可得 (u, v) = (20, 80) cm 或 (80, 20) cm。")
add_equation(doc, "f = uv/(u + v) = 20 × 80/100 = 16 cm。")
add_answer_text(doc, "较大像的放大倍数为 v/u = 80/20 = 4，像高为 4 × 2.0 = 8.0 cm。", indent=True)

doc.core_properties.title = "初三升高一物理竞赛测试题（填空题）"
doc.core_properties.subject = "初三升高一物理竞赛测试，含 TikZ 黑白图转换与参考答案"
doc.core_properties.author = ""
doc.core_properties.keywords = "物理竞赛, 初三升高一, 填空题"

doc.save(OUTPUT)
print(OUTPUT)
