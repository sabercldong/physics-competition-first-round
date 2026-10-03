from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

fonts = [
    "PingFang SC", "Hiragino Sans GB", "STSong", "Songti SC",
    "Arial Unicode MS", "Apple LiSung", "Kaiti SC", "SimSong",
    "Arial", "Noto Sans CJK SC", "Microsoft YaHei", "WenQuanYi Zen Hei",
]

doc = Document()
for font in fonts:
    p = doc.add_paragraph()
    r = p.add_run(f"{font}：中文字体测试——初三升高一物理竞赛")
    r.font.name = font
    r.font.size = Pt(16)
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font)

doc.save("/Users/dongxingxi/Documents/物理竞赛一轮资料/output/tests/初三升高一物理竞赛测试/font_test.docx")
